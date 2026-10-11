from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUT = Path(__file__).parent / 'presentation' / 'RiskLens_Demo.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
NAVY = RGBColor(13, 24, 43); PANEL = RGBColor(25, 43, 69); TEAL = RGBColor(79, 210, 184); WHITE = RGBColor(241, 246, 252); MUTED = RGBColor(177, 194, 215)

def bg(slide):
    fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = NAVY

def textbox(slide, x,y,w,h,text,size=20,color=WHITE,bold=False):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color; p.font.name='Aptos'; p.space_after=Pt(9)
    return box

def base(title, kicker='RISK INTELLIGENCE / MVP'):
    slide=prs.slides.add_slide(prs.slide_layouts[6]); bg(slide)
    textbox(slide,.65,.35,12,.35,kicker,10,TEAL,True)
    textbox(slide,.65,.78,12, .75,title,30,WHITE,True)
    slide.shapes.add_shape(1, Inches(.65), Inches(1.58), Inches(12), Inches(.025)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb=TEAL
    return slide

def card(slide,x,y,w,h,title,body):
    shape=slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h)); shape.fill.solid(); shape.fill.fore_color.rgb=PANEL; shape.line.color.rgb=RGBColor(43,65,94)
    textbox(slide,x+.18,y+.18,w-.36,.4,title,16,TEAL,True)
    textbox(slide,x+.18,y+.68,w-.36,h-.8,body,13,WHITE)

s=base('From market noise to portfolio stress')
textbox(s,.8,2.0,7.2,1.25,'AI-powered Financial Risk Intelligence\n& Portfolio Stress Testing',27,WHITE,True)
textbox(s,.8,4.0,7.4,1.1,'RiskLens converts headlines and social chatter into traceable risk signals, then tests a multi-asset portfolio against configurable shocks.',18,MUTED)
card(s,9.0,2.0,3.3,1.35,'SIGNALS','Sentiment · Event · Impact')
card(s,9.0,3.65,3.3,1.35,'STRESS','Before/after P&L')
textbox(s,.8,6.65,11,.3,'DEMO DATA AND POSITIONS ARE SYNTHETIC · NOT INVESTMENT ADVICE',10,TEAL,True)

s=base('The workflow is designed to be inspectable')
steps=[('01  INGEST','Bundled samples\nRSS · CSV · JSON'),('02  ANALYZE','FinBERT optional\nlexicon fallback'),('03  SIGNAL','Event type\nImpact 1–10'),('04  STRESS','Scenario shocks\nHoldings P&L'),('05  REVIEW','Streamlit + Plotly\nSQLite + exports')]
for i,(t,b) in enumerate(steps):
    x=.7+i*2.52; card(s,x,2.5,2.2,2.1,t,b)
    if i<4: textbox(s,x+2.18,3.25,.35,.4,'→',22,TEAL,True)
textbox(s,.8,5.5,11,.7,'Provenance labels and transparent assumptions keep the demo usable even when external APIs or model downloads fail.',17,MUTED)

s=base('Architecture: local-first, modular, reproducible')
textbox(s,.9,1.95,11.5,.4,'INPUTS',12,TEAL,True)
for i,t in enumerate(['Synthetic\nNewswire','Synthetic\nSocial Pulse','Optional RSS','CSV / JSON']): card(s,.8+i*3.05,2.4,2.6,1.05,t,'')
textbox(s,5.7,3.55,2,.35,'↓ normalize',13,TEAL,True)
card(s,3.4,4.0,6.5,1.05,'NLP RISK ENGINE','Sentiment · Event classification · Impact score · Risk signal')
textbox(s,5.7,5.15,2,.35,'↓ scenario',13,TEAL,True)
card(s,.8,5.6,5.4,1.0,'STRESS ENGINE','Portfolio valuation · Asset-class contributions')
card(s,7.1,5.6,5.4,1.0,'APP + STORAGE','Streamlit · Plotly · SQLite · CSV/JSON exports')

s=base('NLP: useful signals, explicit limitations')
card(s,.8,2.0,3.8,2.35,'SENTIMENT','Signed score in [-1, +1]\nFinBERT optional; fallback is deterministic and local.')
card(s,4.8,2.0,3.8,2.35,'EVENT CLASS','Geopolitical · Macroeconomic\nCredit · M&A · Earnings\nRegulatory · Liquidity · Cyber')
card(s,8.8,2.0,3.8,2.35,'IMPACT 1–10','Keyword-based severity proxy, not a calibrated forecast. Signals map to candidate stress scenarios.')
textbox(s,.9,5.0,11.4,1.0,'Model availability is surfaced in each result. The fallback enables offline demos but is not a substitute for validated financial NLP.',18,MUTED)

s=base('Stress lab: quantify scenario exposure')
card(s,.8,2.0,3.75,2.0,'PORTFOLIO','12 fictional holdings\nEquities · Bonds · Loans · Derivatives')
card(s,4.8,2.0,3.75,2.0,'SCENARIOS','Credit crisis · Rate shock\nGeopolitical · Liquidity · Rebound')
card(s,8.8,2.0,3.75,2.0,'OUTPUTS','Before/after value · P&L\nAsset-class attribution · History')
textbox(s,.9,4.8,11.4,1.1,'Illustrative class-level shocks are multiplied by a user-selected severity. Positions stay static; derivatives use simplified MTM proxies.',18,MUTED)

s=base('Validation, limitations & next steps')
card(s,.8,2.0,3.8,2.7,'VALIDATION','Automated tests cover NLP output bounds, event labels, portfolio valuation, zero-severity invariance and risk-signal mapping.')
card(s,4.8,2.0,3.8,2.7,'LIMITATIONS','No historical calibration, dynamic hedging, transaction costs, loan recovery modeling, derivative Greeks or live-data guarantees.')
card(s,8.8,2.0,3.8,2.7,'ROADMAP','Calibrate shocks on historical data; benchmark NLP; add authenticated data vendors, model monitoring and audit controls.')
textbox(s,.9,5.55,11.4,.6,'DEMO TAKEAWAY  ·  Traceable text signal → configurable scenario → transparent portfolio impact',16,TEAL,True)

prs.save(OUT)
print(OUT)
